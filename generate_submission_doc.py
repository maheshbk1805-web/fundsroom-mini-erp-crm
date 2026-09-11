import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

PRIMARY = RGBColor(16, 42, 34)       # Deep Forest Green #102a22
ACCENT = RGBColor(25, 111, 80)       # Fundsroom Green #196f50
MUTED = RGBColor(101, 117, 109)      # Slate Green #65756d
DARK = RGBColor(23, 32, 29)          # Charcoal #17201d

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = PRIMARY
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    return p

def add_body(text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.color.rgb = DARK
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = DARK
    return p

# TITLE SECTION
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(10)
p_title.paragraph_format.space_after = Pt(2)
r_sub = p_title.add_run("FUNDSROOM RECRUITMENT — ROUND 1 TECHNICAL CASE STUDY\n")
r_sub.font.name = 'Calibri'
r_sub.font.size = Pt(11)
r_sub.font.bold = True
r_sub.font.color.rgb = MUTED

r_main = p_title.add_run("Mini ERP + CRM Operations Portal\nFull-Stack Technical Solution & Architecture Document")
r_main.font.name = 'Calibri'
r_main.font.size = Pt(20)
r_main.font.bold = True
r_main.font.color.rgb = PRIMARY

p_meta = doc.add_paragraph()
p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_meta.paragraph_format.space_after = Pt(16)
r_m = p_meta.add_run("Position: Full Stack Developer | Location: Pune, Maharashtra | Role: Immediate Joining")
r_m.font.name = 'Calibri'
r_m.font.size = Pt(10)
r_m.font.italic = True
r_m.font.color.rgb = MUTED

# DIVIDER
doc.add_paragraph().paragraph_format.space_after = Pt(6)

# 1. EXECUTIVE SUMMARY
add_heading_1("1. Executive Summary & Overview")
add_body("This document presents the comprehensive, production-ready solution for the Fundsroom Full Stack Developer Case Study. The project is an operational Mini ERP and CRM Portal engineered specifically for a wholesale and distribution enterprise. The system coordinates customer relationship lifecycles, catalog inventory with intelligent low-stock triggers, tamper-evident stock movement audit logs, and transactional sales delivery challans.")
add_body("The implementation adheres strictly to the required modern web architecture: a high-performance React + TypeScript SPA on the client, an Express.js + TypeScript RESTful API on the backend, and a robust PostgreSQL database engine ensuring ACID guarantees.")

# 2. TECHNOLOGY STACK
add_heading_1("2. Technical Stack Architecture")
stack_table = doc.add_table(rows=1, cols=3)
stack_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = stack_table.rows[0].cells
hdr[0].text = "Tier / Component"
hdr[1].text = "Technology Chosen"
hdr[2].text = "Justification & Architectural Role"
for cell in hdr:
    set_cell_background(cell, "102A22")
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)

stack_data = [
    ("Frontend Client", "React 18, TypeScript, Vite", "Type safety, instant HMR, modular component structure, and sub-second bundle compilation."),
    ("UI & Styling", "Custom Responsive CSS, Lucide Icons", "Lightweight, ultra-responsive admin dashboard UI adhering to professional enterprise design guidelines."),
    ("Backend API", "Node.js 22, Express.js, TypeScript", "Scalable asynchronous RESTful service with strict request validation and centralized exception handling."),
    ("Database Tier", "PostgreSQL (Neon Cloud & Docker)", "ACID-compliant relational engine with row-level locking (FOR UPDATE) for zero-race-condition transactions."),
    ("Authentication", "JWT (JSON Web Tokens) & Bcrypt", "Stateless, cryptographic authentication with salted password hashing and role-based route middleware."),
    ("DevOps / Containers", "Docker, Docker Compose, Nginx", "Standardized containerization across database, backend API, and static frontend web server."),
    ("CI / CD Automation", "GitHub Actions", "Automated pipeline validating TypeScript compilation, linting, and production builds on push/PR.")
]

for row_data in stack_data:
    row = stack_table.add_row()
    for i, val in enumerate(row_data):
        cell = row.cells[i]
        cell.text = val
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = DARK

# 3. TEST LOGIN CREDENTIALS
add_heading_1("3. Verified Test Login Credentials")
add_body("All accounts are pre-seeded in the database with the uniform password: ", "")
p_pwd = doc.add_paragraph()
r_pwd = p_pwd.add_run("Master Demo Password: Password@123")
r_pwd.font.bold = True
r_pwd.font.color.rgb = ACCENT
r_pwd.font.size = Pt(11)

cred_table = doc.add_table(rows=1, cols=4)
cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
chdr = cred_table.rows[0].cells
chdr[0].text = "Role"
chdr[1].text = "Login Email"
chdr[2].text = "Password"
chdr[3].text = "Module Access & Permissions Scope"
for cell in chdr:
    set_cell_background(cell, "196F50")
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)

roles_data = [
    ("Admin", "admin@fundsroom.local", "Password@123", "Full system access: CRM, Inventory, Stock Movements, Sales Challans, System Analytics."),
    ("Sales", "sales@fundsroom.local", "Password@123", "Sales domain: Create/edit customers, log follow-ups, draft and confirm multi-item sales challans."),
    ("Warehouse", "warehouse@fundsroom.local", "Password@123", "Inventory domain: Create/edit products, manage reorder alerts, record manual stock inward/outward."),
    ("Accounts", "accounts@fundsroom.local", "Password@123", "Auditing domain: Read-only operational visibility, printable tax challans and transaction audit logs.")
]

for row_data in roles_data:
    row = cred_table.add_row()
    for i, val in enumerate(row_data):
        cell = row.cells[i]
        cell.text = val
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = DARK

# 4. CORE MODULE IMPLEMENTATION
add_heading_1("4. Core Modules Implementation Details")

add_heading_2("Module 1: Authentication & Role-Based Access Control (RBAC)")
add_bullet("Secured JWT authentication flow with bearer token header transmission.")
add_bullet("Express middleware 'allow(...roles)' rigorously enforces server-side role boundaries.")
add_bullet("Frontend shell dynamically renders role-specific badges, workspace designations, and contextual action controls.")

add_heading_2("Module 2: Customer CRM Management")
add_bullet("Fields tracked: Customer Name, Mobile, Email, Business Name, GST Number (Optional), Customer Type (Retail, Wholesale, Distributor), Address, Status (Lead, Active, Inactive), Follow-up Date, and Notes.")
add_bullet("Real-time search functionality querying customer names and business entities.")
add_bullet("Dedicated Customer Profile Drawer displaying contact records, registered GSTIN, and full history.")
add_bullet("Interactive Follow-up Log: Sales users log chronological interactions and dynamically reschedule the next action date.")

add_heading_2("Module 3: Products & Inventory Management")
add_bullet("Fields tracked: Product Name, SKU Code, Category, Unit Price (INR), Current On-Hand Stock, Minimum Stock Alert Threshold, and Warehouse Location.")
add_bullet("Dynamic low-stock indicators displaying red warning tags when on-hand quantity <= minimum threshold.")
add_bullet("In-place modal editing allowing warehouse managers to update price, stock levels, and alert rules.")

add_heading_2("Module 4: Stock Movement Audit Trail")
add_bullet("Immutable record logging: Product, SKU, Movement Type (IN / OUT), Quantity Changed, Business Reason, Authorized User, and Timestamp.")
add_bullet("Automatic background creation of 'OUT' movements upon challan confirmation, providing end-to-end trace integrity.")

add_heading_2("Module 5: Sales Challan Flow & Concurrency Safety")
add_bullet("Multi-Product Line Items: Sales users can dynamically add multiple product items to a single challan with live subtotal and inventory checks.")
add_bullet("Automatic Challan Number Generation in formatted sequence (e.g. CH-2026-XXXXXX).")
add_bullet("Draft vs. Confirmed status: Drafts reserve no inventory; Confirmed status atomically decrements inventory.")
add_bullet("Historical Pricing Snapshot: Challan line items preserve product name, SKU, and unit price in challan_items to guarantee invoice immutability.")
add_bullet("Row-Level Locking (SELECT ... FOR UPDATE) prevents race conditions and overselling.")
add_bullet("Zero Negative Stock Enforcement: If available inventory is insufficient, the transaction rolls back with HTTP 409 Conflict.")

# 5. BONUS FEATURES IMPLEMENTATION
add_heading_1("5. Bonus Points Implementation")
add_bullet("Docker Setup: Complete docker-compose.yml orchestrating PostgreSQL, multi-stage Node backend, and Nginx frontend.", "1. Containerization: ")
add_bullet("GitHub Actions CI: Fully automated .github/workflows/ci.yml pipeline for continuous integration and typechecking.", "2. Automated CI/CD: ")
add_bullet("Invoice & Delivery Challan PDF Export: One-click print/PDF layout with company headers, customer GST, itemized totals, and signatures.", "3. PDF Export: ")

# 6. HOW TO RUN LOCALLY & VERIFICATION
add_heading_1("6. Step-by-Step Local Execution Guide")
add_body("The project is 100% verified and operational on the development workstation.")
add_body("Step 1: Open VS Code terminal in the project directory (fundsroom project).")
add_body("Step 2: Start both client and server concurrently by executing:")
p_cmd = doc.add_paragraph()
r_c = p_cmd.add_run("pnpm dev   (or npm run dev)")
r_c.font.name = 'Consolas'
r_c.font.bold = True
r_c.font.color.rgb = ACCENT
add_body("Step 3: Open your browser and navigate to: http://localhost:5173")
add_body("Step 4: Log in using any of the 4 test accounts listed in Section 3.")

# 7. SCREEN RECORDING & SUBMISSION CHECKLIST
add_heading_1("7. Screen Recording & Submission Checklist")
add_body("As mandated in the assessment instructions, ensure the following steps are performed:")
add_bullet("Start your screen recorder (OBS Studio, Windows Game Bar [Win + G], or Loom).", "1. Screen Recording: ")
add_bullet("Demonstrate Admin login, overview metrics, and customer creation with GST number.", "2. CRM Flow: ")
add_bullet("Demonstrate customer detail view and logging a follow-up note with new date.", "3. Follow-up Flow: ")
add_bullet("Demonstrate product catalog, low stock indicator, and editing a product.", "4. Inventory Flow: ")
add_bullet("Create a multi-product Sales Challan as Draft, view its details, and confirm it.", "5. Challan Flow: ")
add_bullet("Demonstrate that stock was decremented and an OUT movement was created in the Audit log.", "6. Stock Verification: ")
add_bullet("Click 'Print / Export PDF' to demonstrate the delivery note invoice.", "7. PDF Export: ")
add_bullet("Upload recording to Google Drive (set permission to 'Anyone with the link can view').", "8. Drive Upload: ")
add_bullet("Submit the Google Form before 6:00 PM tomorrow with PDF solution & Drive link.", "9. Form Submission: ")

# 8. CONCLUSION & JOINING CONFIRMATION
add_heading_1("8. Immediate Joining Confirmation")
add_body("I confirm my availability for immediate joining at Fundsroom in Pune, Maharashtra. I am fully comfortable with the Pune work location and look forward to contributing to Fundsroom's technology roadmap.")

output_path = "Fundsroom_Full_Stack_Case_Study_Solution.docx"
doc.save(output_path)
print(f"Document successfully created: {output_path}")
